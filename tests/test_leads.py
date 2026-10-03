import json
import threading
import urllib.request
from datetime import datetime
from http.server import HTTPServer

import pytest
from openpyxl import load_workbook

from lead_automation.classify import keyword_intent, score
from lead_automation.generate_samples import generate
from lead_automation.models import Lead
from lead_automation.pipeline import LeadPipeline
from lead_automation.report import write_crm_csv, write_excel
from lead_automation.sources import from_csv, from_email_folder, from_json_folder
from lead_automation.webhook import make_handler


@pytest.fixture(autouse=True)
def no_llm(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)


@pytest.mark.parametrize("text, intent", [
    ("We offer SEO services and backlinks", "spam"),
    ("Adjunto mi curriculum para la vacante", "job_application"),
    ("My order is not working, I want a refund", "support"),
    ("Interested in a reseller partnership", "partnership"),
    ("Can you send me pricing for 20 users?", "sales"),
])
def test_keyword_intent(text, intent):
    assert keyword_intent(text) == intent


def test_hot_lead_scoring_explains_reasons():
    lead = Lead(name="ana gil", email="ana@acme.com", phone="+5491111111111", company="Acme", source="website",
                budget=8000, employees=120, message="Need a quote ASAP")
    lead.intent = "sales"
    score(lead)
    assert lead.tier == "Hot" and lead.score >= 70
    assert any("Budget" in r for r in lead.reasons)


def test_cold_and_disqualified():
    cold = Lead(name="x", email="x@gmail.com", message="just browsing", source="ads_csv")
    cold.intent = "sales"
    score(cold)
    assert cold.tier == "Cold"
    spam = Lead(name="seo", email="s@seo.biz", message="backlinks")
    spam.intent = "spam"
    score(spam)
    assert spam.tier == "Disqualified" and spam.score == 0


def test_name_and_email_normalization():
    lead = Lead(name="  laura   MÉNDEZ ", email=" Laura@Ferreteria.COM ")
    assert lead.name == "Laura Méndez"
    assert lead.email == "laura@ferreteria.com"


@pytest.fixture
def samples(tmp_path):
    generate(tmp_path)
    return tmp_path


def test_sources_and_dedupe(samples):
    leads = from_csv(samples / "facebook_lead_ads.csv") + from_json_folder(samples / "web_forms") \
        + from_email_folder(samples / "emails")
    pipeline = LeadPipeline(use_llm=False)
    stats = pipeline.process_many(leads)
    assert stats["received"] == 7 + 5 + 4
    assert stats["duplicates"] == 2  # two people filled both the ad form and the website form
    rows = pipeline.all_leads()
    assert {r["intent"] for r in rows} >= {"spam", "job_application", "support", "sales"}
    assert rows[0]["score"] >= rows[-1]["score"]  # sorted by score
    assert {r["owner"] for r in rows if r["intent"] == "support"} == {"support@"}


def test_email_parsing_extracts_company_budget_phone(samples):
    lead = [l for l in from_email_folder(samples / "emails") if "clinicadelsur" in l.email][0]
    assert lead.company == "Clínica del Sur"
    assert lead.budget == 9000
    assert lead.phone.startswith("+54")


def test_reports(samples, tmp_path):
    pipeline = LeadPipeline(use_llm=False)
    pipeline.process_many(from_csv(samples / "facebook_lead_ads.csv"))
    rows = pipeline.all_leads()
    due = pipeline.due_follow_ups(datetime(2030, 1, 1))
    path = write_excel(rows, due, tmp_path / "r.xlsx")
    assert load_workbook(path).sheetnames == ["Summary", "Pipeline", "Follow-ups due"]
    csv_text = write_crm_csv(rows, tmp_path / "crm.csv").read_text(encoding="utf-8")
    assert csv_text.startswith("First Name,Last Name,Email")


def test_webhook_creates_and_dedupes():
    pipeline = LeadPipeline(use_llm=False)
    server = HTTPServer(("127.0.0.1", 0), make_handler(pipeline))
    threading.Thread(target=server.serve_forever, daemon=True).start()
    url = f"http://127.0.0.1:{server.server_port}/lead"
    body = json.dumps({"name": "Ana Gil", "email": "ana@acme.com", "message": "quote asap", "budget": 8000}).encode()

    def post():
        req = urllib.request.Request(url, data=body, headers={"Content-Type": "application/json"})
        with urllib.request.urlopen(req) as resp:
            return resp.status, json.loads(resp.read())

    first, second = post(), post()
    server.shutdown()
    assert first[0] == 201 and first[1]["status"] == "created"
    assert second[0] == 200 and second[1]["status"] == "duplicate"
