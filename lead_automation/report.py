"""Excel pipeline report + CRM-ready CSV."""

from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

TIER_COLORS = {"Hot": "F8CBAD", "Warm": "FFE699", "Cold": "DDEBF7", "Disqualified": "E7E6E6"}
HEADER_FILL = PatternFill("solid", fgColor="7030A0")
HEADER_FONT = Font(bold=True, color="FFFFFF")
COLUMNS = ["score", "tier", "intent", "name", "company", "email", "phone", "source", "budget", "owner",
           "next_action", "follow_up", "times_seen", "reasons", "message"]


def _sheet(wb: Workbook, title: str, rows: list[dict], first: bool = False) -> Worksheet:
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append([c.replace("_", " ").title() for c in COLUMNS])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for row in rows:
        reasons = "; ".join(json.loads(row["reasons"])) if row.get("reasons") else ""
        ws.append([reasons if c == "reasons" else row.get(c) for c in COLUMNS])
        fill = TIER_COLORS.get(row.get("tier"))
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=fill)
    widths = {"message": 60, "reasons": 60, "next_action": 30, "email": 30}
    for idx, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(col, 16)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def write_excel(leads: list[dict], due: list[dict], output: str | Path) -> Path:
    """Write the Summary / Pipeline / Follow-ups due workbook and return its path."""
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _sheet(wb, "Pipeline", leads, first=True)
    _sheet(wb, "Follow-ups due", due)

    ws = wb.create_sheet("Summary", 0)
    ws.append(["Lead pipeline summary"])
    ws["A1"].font = Font(bold=True, size=14)
    for title, counter in (("Tier", Counter(r["tier"] for r in leads)),
                           ("Source", Counter(r["source"] for r in leads)),
                           ("Owner", Counter(r["owner"] for r in leads))):
        ws.append([])
        ws.append([title, "Leads"])
        for cell in ws[ws.max_row]:
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        for key, count in counter.most_common():
            ws.append([key, count])
    ws.column_dimensions["A"].width = 24
    wb.save(output)
    return output


def write_crm_csv(leads: list[dict], output: str | Path) -> Path:
    """Columns match a typical CRM import (HubSpot / Pipedrive style). Disqualified leads are excluded."""
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    fields = ["First Name", "Last Name", "Email", "Phone Number", "Company Name", "Lead Status", "Lead Score",
              "Contact Owner", "Original Source"]
    with open(output, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        for r in leads:
            if r["tier"] == "Disqualified":
                continue
            first, _, last = (r["name"] or "").partition(" ")
            writer.writerow({"First Name": first, "Last Name": last, "Email": r["email"], "Phone Number": r["phone"],
                             "Company Name": r["company"], "Lead Status": r["tier"], "Lead Score": r["score"],
                             "Contact Owner": r["owner"], "Original Source": r["source"]})
    return output
